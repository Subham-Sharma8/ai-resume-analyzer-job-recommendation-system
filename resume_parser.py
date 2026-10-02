import io
from pathlib import Path
from pypdf import PdfReader
from docx import Document

ALLOWED_EXTENSIONS = {".pdf", ".docx"}

def extract_resume_text(file_bytes: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file type. Upload PDF or DOCX.")

    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(file_bytes))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return "\n".join(pages).strip()

    document = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in document.paragraphs]
    table_text = []
    for table in document.tables:
        for row in table.rows:
            table_text.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs + table_text).strip()
